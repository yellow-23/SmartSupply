import math
import re
from datetime import date, timedelta
from io import BytesIO
from typing import Annotated, Optional

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from fpdf import FPDF
from openpyxl import Workbook
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.api.auth import assert_business_access, get_current_user
from app.database import get_db
from app.models.orm import Business, Product, PurchaseOrder, PurchaseOrderLine, StockLevel, Store, User
from app.models.schemas import PurchaseOrderCreate, PurchaseOrderLineIn, PurchaseOrderResponse, PurchaseOrderUpdate
from app.services.inventory_service import InventoryService, _get_product_params

router = APIRouter()
service = InventoryService()

_VALID_STATUSES = {"pending", "confirmed", "in_transit", "received", "cancelled"}
_TRANSITIONS = {
    "pending": {"confirmed", "cancelled"},
    "confirmed": {"in_transit", "cancelled"},
    "in_transit": {"received", "cancelled"},
    "received": set(),
    "cancelled": set(),
}


def _filtered_orders(
    db: Session,
    business_id: int,
    store_nbr: Optional[int],
    status: Optional[str],
    q: Optional[str],
    date_from: Optional[date],
    date_to: Optional[date],
):
    query = db.query(PurchaseOrder).filter(PurchaseOrder.business_id == business_id)
    if store_nbr is not None:
        query = query.filter(PurchaseOrder.store_nbr == store_nbr)
    if status:
        query = query.filter(PurchaseOrder.status == status)
    if q:
        q = q.strip()
        number = re.fullmatch(r"(?i)(?:oc-?)?0*(\d+)", q)
        like = f"%{q}%"
        cond = PurchaseOrder.supplier.ilike(like) | PurchaseOrder.lines.any(PurchaseOrderLine.family.ilike(like))
        if number:
            cond = cond | (PurchaseOrder.id == int(number.group(1)))
        query = query.filter(cond)
    if date_from:
        query = query.filter(func.date(PurchaseOrder.created_at) >= date_from)
    if date_to:
        query = query.filter(func.date(PurchaseOrder.created_at) <= date_to)
    return query.order_by(PurchaseOrder.created_at.desc(), PurchaseOrder.id.desc())


@router.get("", response_model=list[PurchaseOrderResponse])
def list_orders(
    business_id: int,
    store_nbr: Optional[int] = None,
    status: Optional[str] = None,
    q: Optional[str] = None,
    date_from: Optional[date] = None,
    date_to: Optional[date] = None,
    current_user: Annotated[User, Depends(get_current_user)] = None,
    db: Session = Depends(get_db),
):
    """Lista las OC del negocio. `q` busca por numero (OC-000123 o 123), producto o proveedor."""
    assert_business_access(db, current_user, business_id)
    return _filtered_orders(db, business_id, store_nbr, status, q, date_from, date_to).all()


@router.get("/export")
def export_orders(
    business_id: int,
    store_nbr: Optional[int] = None,
    status: Optional[str] = None,
    q: Optional[str] = None,
    date_from: Optional[date] = None,
    date_to: Optional[date] = None,
    current_user: Annotated[User, Depends(get_current_user)] = None,
    db: Session = Depends(get_db),
):
    """Exporta las OC (mismos filtros que /orders) a Excel, una fila por linea de producto."""
    assert_business_access(db, current_user, business_id)
    orders = _filtered_orders(db, business_id, store_nbr, status, q, date_from, date_to).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Ordenes"
    ws.append([
        "N° OC", "Proveedor", "Tienda", "Estado", "Producto", "Cantidad", "Costo unit.", "Subtotal",
        "Stock al disparar", "Punto s", "Nivel S", "Creada", "Entrega esperada", "Recibida",
    ])
    num = lambda v: float(v) if v is not None else None
    for o in orders:
        for l in o.lines:
            ws.append([
                f"OC-{o.id:06d}",
                o.supplier,
                o.store_nbr,
                o.status,
                l.family,
                num(l.quantity),
                num(l.unit_cost),
                float(l.quantity * l.unit_cost) if l.unit_cost is not None else None,
                num(l.trigger_stock),
                num(l.reorder_point_s),
                num(l.order_up_to_S),
                o.created_at.strftime("%Y-%m-%d %H:%M") if o.created_at else None,
                o.expected_delivery.strftime("%Y-%m-%d") if o.expected_delivery else None,
                o.received_at.strftime("%Y-%m-%d %H:%M") if o.received_at else None,
            ])
    for col in "ABCDEFGHIJKLMN":
        ws.column_dimensions[col].width = 16

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=ordenes_negocio{business_id}.xlsx"},
    )


@router.post("/generate", response_model=list[PurchaseOrderResponse])
def generate_automatic_orders(
    business_id: int,
    store_nbr: int = 1,
    current_user: Annotated[User, Depends(get_current_user)] = None,
    db: Session = Depends(get_db),
):
    """
    Revisa todos los SKUs en estado crítico (stock <= s) y genera órdenes
    automáticas con la política (s,S). Devuelve las órdenes creadas.
    """
    assert_business_access(db, current_user, business_id)
    try:
        critical_skus = service.get_critical_skus(db=db, business_id=business_id, store_nbr=store_nbr)
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

    open_families = {
        f for (f,) in db.query(PurchaseOrderLine.family).join(PurchaseOrder).filter(
            PurchaseOrder.business_id == business_id,
            PurchaseOrder.store_nbr == store_nbr,
            PurchaseOrder.status.in_(["pending", "confirmed", "in_transit"]),
        ).distinct()
    }

    # Una OC por proveedor; los productos sin proveedor asignado van juntos en una OC aparte.
    by_supplier: dict[Optional[str], list] = {}
    for sku in critical_skus:
        if sku["family"] in open_families:
            # Ya hay una orden abierta que cubre este SKU; generar otra la duplicaria.
            continue
        if sku.get("needs_cost_setup"):
            # Sin unit_cost configurado el calculo de cantidad no es confiable -- no se genera
            # orden automatica para este SKU hasta que se cargue el costo (ver /inventory/alerts).
            continue
        params = _get_product_params(db, business_id, store_nbr, sku["family"])
        by_supplier.setdefault(params.get("supplier"), []).append((sku, params))

    created = []
    for supplier, items in by_supplier.items():
        order = PurchaseOrder(
            business_id=business_id,
            store_nbr=store_nbr,
            supplier=supplier,
            policy_used="s_s",
            status="pending",
            expected_delivery=date.today() + timedelta(days=max(p["lead_time_days"] for _, p in items)),
            lines=[
                PurchaseOrderLine(
                    family=sku["family"],
                    quantity=math.ceil(sku["order_quantity"]),
                    unit_cost=params.get("unit_cost"),
                    trigger_stock=round(sku["current_stock"]),
                    reorder_point_s=round(sku["reorder_point_s"]),
                    order_up_to_S=round(sku["order_up_to_S"]),
                )
                for sku, params in items
            ],
        )
        db.add(order)
        created.append(order)

    db.commit()
    for o in created:
        db.refresh(o)

    return created


def _build_lines(db: Session, business_id: int, store_nbr: int, lines: list[PurchaseOrderLineIn]):
    """Valida que cada producto exista en el catalogo y completa el costo desde ahi si no viene.
    Devuelve (lineas, lead time maximo)."""
    families = [l.family.strip().upper() for l in lines]
    if len(set(families)) != len(families):
        raise HTTPException(status_code=400, detail="Hay productos repetidos en la orden.")
    products = {
        p.family: p for p in db.query(Product).filter(
            Product.business_id == business_id, Product.store_nbr == store_nbr, Product.family.in_(families),
        )
    }
    missing = [f for f in families if f not in products]
    if missing:
        raise HTTPException(status_code=400, detail=f"Productos no registrados en el catalogo: {', '.join(missing)}")
    built = [
        PurchaseOrderLine(
            family=f,
            quantity=l.quantity,
            unit_cost=l.unit_cost if l.unit_cost is not None else products[f].unit_cost,
        )
        for f, l in zip(families, lines)
    ]
    lead = max((products[f].lead_time_days or 7) for f in families)
    return built, lead


def _get_order(db: Session, order_id: int, business_id: int) -> PurchaseOrder:
    order = db.query(PurchaseOrder).filter(
        PurchaseOrder.id == order_id,
        PurchaseOrder.business_id == business_id,
    ).first()
    if not order:
        raise HTTPException(status_code=404, detail=f"Orden #{order_id} no encontrada")
    return order


@router.post("", response_model=PurchaseOrderResponse, status_code=201)
def create_order(
    body: PurchaseOrderCreate,
    current_user: Annotated[User, Depends(get_current_user)] = None,
    db: Session = Depends(get_db),
):
    """Crea una OC manual. El costo de cada linea se toma del catalogo si no se indica."""
    assert_business_access(db, current_user, body.business_id)
    lines, lead = _build_lines(db, body.business_id, body.store_nbr, body.lines)
    order = PurchaseOrder(
        business_id=body.business_id,
        store_nbr=body.store_nbr,
        supplier=(body.supplier or "").strip() or None,
        policy_used="manual",
        status="pending",
        expected_delivery=body.expected_delivery or date.today() + timedelta(days=lead),
        lines=lines,
    )
    db.add(order)
    db.commit()
    db.refresh(order)
    return order


@router.patch("/{order_id}", response_model=PurchaseOrderResponse)
def update_order(
    order_id: int,
    business_id: int,
    body: PurchaseOrderUpdate,
    current_user: Annotated[User, Depends(get_current_user)] = None,
    db: Session = Depends(get_db),
):
    """Edita proveedor, fecha de entrega o lineas. Solo mientras la OC esta pendiente."""
    assert_business_access(db, current_user, business_id)
    order = _get_order(db, order_id, business_id)
    if order.status != "pending":
        raise HTTPException(status_code=400, detail="Solo se puede editar una orden pendiente.")
    fields = body.model_dump(exclude_unset=True)
    if "supplier" in fields:
        order.supplier = (body.supplier or "").strip() or None
    if "expected_delivery" in fields:
        order.expected_delivery = body.expected_delivery
    if body.lines is not None:
        order.lines, _ = _build_lines(db, business_id, order.store_nbr, body.lines)
    db.commit()
    db.refresh(order)
    return order


@router.get("/{order_id}/pdf")
def export_order_pdf(
    order_id: int,
    business_id: int,
    current_user: Annotated[User, Depends(get_current_user)] = None,
    db: Session = Depends(get_db),
):
    """Documento de la OC para enviar al proveedor."""
    assert_business_access(db, current_user, business_id)
    order = _get_order(db, order_id, business_id)
    data = PurchaseOrderResponse.model_validate(order)
    business = db.query(Business).filter(Business.id == business_id).first()
    store = db.query(Store).filter(Store.business_id == business_id, Store.store_nbr == data.store_nbr).first()
    clp = lambda v: f"${v:,.0f}".replace(",", ".")
    # fpdf con fuentes base solo soporta latin-1
    txt = lambda v: str(v).encode("latin-1", "replace").decode("latin-1")

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 18)
    pdf.cell(0, 10, f"Orden de Compra {data.order_number}", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 11)
    buyer_contact = " - ".join(x for x in (business.phone, business.email) if x) if business else ""
    buyer_address = ", ".join(x for x in (business.address, business.city) if x) if business else ""
    delivery = (store.name if store and store.name else f"Tienda {data.store_nbr}") + (f", {store.city}" if store and store.city else "")
    rows = [
        ("Comprador", business.name if business else str(business_id)),
        ("RUT", business.rut if business and business.rut else "-"),
        ("Direccion", buyer_address or "-"),
        ("Contacto", buyer_contact or "-"),
        ("Entregar en", delivery),
        ("Proveedor", data.supplier or "Sin proveedor asignado"),
        ("Fecha emision", data.created_at.strftime("%d/%m/%Y")),
        ("Entrega esperada", data.expected_delivery.strftime("%d/%m/%Y") if data.expected_delivery else "-"),
        ("Estado", {"pending": "Pendiente", "confirmed": "Confirmada", "in_transit": "En transito",
                    "received": "Recibida", "cancelled": "Cancelada"}.get(data.status, data.status)),
    ]
    for label, value in rows:
        pdf.set_font("Helvetica", "B", 11)
        pdf.cell(45, 7, label)
        pdf.set_font("Helvetica", "", 11)
        pdf.cell(0, 7, txt(value), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(6)

    pdf.set_font("Helvetica", "", 10)
    with pdf.table(col_widths=(80, 30, 35, 35), text_align=("LEFT", "RIGHT", "RIGHT", "RIGHT")) as table:
        header = table.row()
        for col in ("Producto", "Cantidad", "Costo unit.", "Subtotal"):
            header.cell(col)
        for l in data.lines:
            row = table.row()
            row.cell(txt(l.family))
            row.cell(f"{l.quantity:,.0f}".replace(",", "."))
            row.cell(clp(l.unit_cost) if l.unit_cost is not None else "-")
            row.cell(clp(l.quantity * l.unit_cost) if l.unit_cost is not None else "-")
    pdf.ln(4)
    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 8, f"Total: {clp(data.total_cost) if data.total_cost is not None else 'pendiente de costos'}",
             align="R", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "I", 9)
    pdf.cell(0, 6, "Generada por SmartSupply", new_x="LMARGIN", new_y="NEXT")

    buf = BytesIO(bytes(pdf.output()))
    return StreamingResponse(
        buf,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={data.order_number}.pdf"},
    )


@router.patch("/{order_id}/status", response_model=PurchaseOrderResponse)
def update_order_status(
    order_id: int,
    new_status: str,
    business_id: int,
    current_user: Annotated[User, Depends(get_current_user)] = None,
    db: Session = Depends(get_db),
):
    """
    Cambia el estado de una orden. Transiciones válidas:
    pending -> confirmed | cancelled
    confirmed -> in_transit | cancelled
    in_transit -> received | cancelled
    """
    if new_status not in _VALID_STATUSES:
        raise HTTPException(status_code=400, detail=f"Estado inválido: {new_status}")
    assert_business_access(db, current_user, business_id)
    order = _get_order(db, order_id, business_id)

    allowed = _TRANSITIONS.get(order.status, set())
    if new_status not in allowed:
        raise HTTPException(
            status_code=400,
            detail=f"No se puede pasar de '{order.status}' a '{new_status}'",
        )

    order.status = new_status
    if new_status == "received":
        from datetime import datetime
        order.received_at = datetime.now()

        for line in order.lines:
            stock = db.query(StockLevel).filter(
                StockLevel.business_id == business_id,
                StockLevel.store_nbr == order.store_nbr,
                StockLevel.family == line.family,
            ).first()
            if stock:
                stock.quantity = float(stock.quantity or 0) + float(line.quantity)
            else:
                db.add(StockLevel(
                    business_id=business_id, store_nbr=order.store_nbr,
                    family=line.family, quantity=line.quantity,
                ))

    db.commit()
    db.refresh(order)
    return order


@router.get("/{order_id}", response_model=PurchaseOrderResponse)
def get_order(
    order_id: int,
    business_id: int,
    current_user: Annotated[User, Depends(get_current_user)] = None,
    db: Session = Depends(get_db),
):
    """Detalle de una orden de compra por ID."""
    assert_business_access(db, current_user, business_id)
    return _get_order(db, order_id, business_id)
