"""Archivo descargable (PDF o Excel) de una prediccion, por producto o del total de una ubicacion."""
from io import BytesIO

from fastapi.responses import StreamingResponse
from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.models.schemas import ForecastResponse

MODEL_NAMES = {"arima": "ARIMA", "prophet": "Prophet", "xgboost": "XGBoost", "lstm": "LSTM"}
DIAS = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo"]


def _num(v: float, clp: bool) -> str:
    s = f"{v:,.0f}".replace(",", ".")
    return f"${s}" if clp else s


def _latin1(v) -> str:
    # fpdf con fuentes base solo soporta latin-1
    return str(v).encode("latin-1", "replace").decode("latin-1")


def _slug(text: str) -> str:
    return "".join(c if c.isalnum() else "_" for c in text).strip("_").lower() or "prediccion"


def forecast_file(result: ForecastResponse, subject: str, business_name: str, store_label: str, fmt: str) -> StreamingResponse:
    """subject: el producto o "Venta total de la tienda". fmt: 'pdf' | 'xlsx'."""
    clp = result.sales_unit == "CLP"
    unit = "CLP" if clp else "unidades"
    preds = result.predictions
    total = sum(p.predicted_sales for p in preds)
    summary = [
        ("Negocio", business_name),
        ("Ubicación", store_label),
        ("Producto", subject),
        ("Período", f"{preds[0].date:%d/%m/%Y} al {preds[-1].date:%d/%m/%Y} ({len(preds)} días)" if preds else "-"),
        ("Modelo elegido", MODEL_NAMES.get(result.model_used, result.model_used)),
        ("Error del modelo (WAPE)", f"{result.mape_validation:.1f}%" if result.mape_validation is not None else "-"),
        (f"Total estimado ({unit})", _num(total, clp)),
        (f"Promedio diario ({unit})", _num(total / max(len(preds), 1), clp)),
        ("Generada", result.generated_at.strftime("%d/%m/%Y %H:%M")),
    ]
    name = f"prediccion_{_slug(subject)}_{preds[0].date:%Y%m%d}" if preds else "prediccion"

    if fmt == "xlsx":
        wb = Workbook()
        ws = wb.active
        ws.title = "Prediccion"
        ws.append(["SmartSupply - Predicción de demanda"])
        ws["A1"].font = Font(bold=True, size=14)
        ws.append([])
        for label, value in summary:
            ws.append([label, value])
            ws.cell(ws.max_row, 1).font = Font(bold=True)
        ws.append([])
        has_bounds = any(p.lower_bound is not None for p in preds)
        header = ["Fecha", "Día", f"Predicción ({unit})"] + (["Mínimo", "Máximo"] if has_bounds else [])
        ws.append(header)
        for c in ws[ws.max_row]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="0F172A")
        money = '"$"#,##0' if clp else "#,##0"
        for p in preds:
            row = [p.date, DIAS[p.date.weekday()], round(p.predicted_sales)]
            if has_bounds:
                row += [p.lower_bound, p.upper_bound]
            ws.append(row)
            ws.cell(ws.max_row, 1).number_format = "DD/MM/YYYY"
            for col in range(3, len(row) + 1):
                ws.cell(ws.max_row, col).number_format = money
        ws.append(["Total", "", round(total)])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        ws.cell(ws.max_row, 3).font = Font(bold=True)
        ws.cell(ws.max_row, 3).number_format = money
        for col, width in zip("ABCDE", (26, 40, 20, 14, 14)):
            ws.column_dimensions[col].width = width
        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        media, ext = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"
    else:
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Helvetica", "B", 16)
        pdf.cell(0, 10, _latin1("Predicción de demanda"), new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 10)
        pdf.cell(0, 6, "SmartSupply", new_x="LMARGIN", new_y="NEXT")
        pdf.ln(3)
        for label, value in summary:
            pdf.set_font("Helvetica", "B", 10)
            pdf.cell(55, 6.5, _latin1(label))
            pdf.set_font("Helvetica", "", 10)
            pdf.cell(0, 6.5, _latin1(value), new_x="LMARGIN", new_y="NEXT")
        pdf.ln(5)
        pdf.set_font("Helvetica", "", 10)
        with pdf.table(col_widths=(40, 40, 50), text_align=("LEFT", "LEFT", "RIGHT")) as table:
            head = table.row()
            for col in ("Fecha", _latin1("Día"), _latin1(f"Predicción ({unit})")):
                head.cell(col)
            for p in preds:
                row = table.row()
                row.cell(p.date.strftime("%d/%m/%Y"))
                row.cell(_latin1(DIAS[p.date.weekday()]))
                row.cell(_num(p.predicted_sales, clp))
        pdf.ln(4)
        pdf.set_font("Helvetica", "I", 8)
        pdf.multi_cell(0, 4.5, _latin1(
            "La predicción es el valor más probable para cada día según el historial cargado. "
            "El WAPE indica cuánto se equivocó el modelo al probarse con datos que no vio al entrenar."
        ))
        buf = BytesIO(bytes(pdf.output()))
        media, ext = "application/pdf", "pdf"

    return StreamingResponse(buf, media_type=media, headers={"Content-Disposition": f"attachment; filename={name}.{ext}"})
