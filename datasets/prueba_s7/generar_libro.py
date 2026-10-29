"""Genera SmartSupply_libro_prueba.xlsx: libro de distribuidora con Instrucciones, Catalogo, Ventas y Stock.
El stock se fija contra el punto s real del backend para que el resultado esperado sea exacto.
Uso (desde la raiz del repo): python3.11 datasets/prueba_s7/generar_libro.py"""
import os
import sys
from datetime import date, timedelta

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
sys.path[:0] = [ROOT, os.path.join(ROOT, "backend")]
os.environ.setdefault("ANTHROPIC_API_KEY", "x")
from app.services.inventory_service import _calc_s_S  # noqa: E402

OUT = os.path.join(os.path.dirname(__file__), "SmartSupply_libro_prueba.xlsx")
rng = np.random.default_rng(23)

# producto: familia, proveedor, costo, precio venta, lead time, moq, venta diaria base, estado de stock deseado
CATALOGO = [
    ("COCA-COLA 1.5L",          "BEBIDAS",   "Embotelladora Andina", 1150, 1890, 3, 24, 42, "critico"),
    ("COCA-COLA ZERO 1.5L",     "BEBIDAS",   "Embotelladora Andina", 1150, 1890, 3, 24, 20, "critico"),
    ("SPRITE 1.5L",             "BEBIDAS",   "Embotelladora Andina", 1050, 1790, 3, 24, 16, "ok"),
    ("FANTA 1.5L",              "BEBIDAS",   "Embotelladora Andina", 1050, 1790, 3, 24, 13, "critico"),
    ("PEPSI 1.5L",              "BEBIDAS",   "CCU",                  990, 1690, 5, 12, 21, "critico"),
    ("BILZ 1.5L",               "BEBIDAS",   "CCU",                  950, 1590, 5, 12, 12, "ok"),
    ("AGUA CACHANTUN 1.6L",     "BEBIDAS",   "CCU",                  690, 1190, 5, 12, 18, "ok"),
    ("LECHE ENTERA SOPROLE 1L", "LACTEOS",   "Soprole",              890, 1290, 2, 12, 36, "critico"),
    ("YOGHURT SOPROLE 165G",    "LACTEOS",   "Soprole",              320, 490, 2, 24, 30, "ok"),
    ("QUESO GAUDA COLUN 250G",  "LACTEOS",   "Colun",                2490, 3490, 4, 6, 9, "critico"),
    ("MANTEQUILLA COLUN 250G",  "LACTEOS",   "Colun",                2190, 2990, 4, 6, 7, "ok"),
    ("ARROZ TUCAPEL 1KG",       "ABARROTES", "Distribuidora Central", 1390, 1990, 4, 10, 15, "critico"),
    ("ACEITE CHEF 1L",          "ABARROTES", "Distribuidora Central", 2290, 3190, 4, 6, 10, "critico"),
    ("AZUCAR IANSA 1KG",        "ABARROTES", "Distribuidora Central", 1090, 1490, 4, 10, 14, "ok"),
    ("FIDEOS CAROZZI 400G",     "ABARROTES", "Carozzi",              690, 990, 3, 20, 22, "critico"),
    ("MARRAQUETA KG",           "PANADERIA", None,                   1600, 2490, 1, None, 26, "critico"),
    ("HALLULLA KG",             "PANADERIA", None,                   1700, 2590, 1, None, 15, "ok"),
    ("DETERGENTE OMO 3KG",      "LIMPIEZA",  "Unilever",             7990, 10990, 7, 4, 3, "critico"),
    ("CLORO CLORINDA 1L",       "LIMPIEZA",  "Comercial Limpia",     790, 1190, 5, 12, 8, "ok"),
    ("PAPEL HIGIENICO ELITE 4U", "LIMPIEZA", "Comercial Limpia",     1990, 2890, 5, 10, 11, "critico"),
    ("PAPAS LAYS 230G",         "SNACKS",    "Evercrisp",            1690, 2490, 6, 12, 9, "ok"),
    ("GALLETAS TRITON 126G",    "SNACKS",    "Nestle",               None, 890, 6, 24, 17, "critico"),  # sin costo a proposito
]
DOW = [0.9, 0.85, 0.9, 1.0, 1.2, 1.45, 0.7]  # lun..dom
FIESTAS = {"BEBIDAS": 2.2, "SNACKS": 2.0, "PANADERIA": 1.6}  # semana del 18 de septiembre

end = date.today() - timedelta(days=1)
days = [end - timedelta(days=i) for i in range(179, -1, -1)]

ventas, stock, esperado = [], [], []
for prod, fam, prov, costo, precio, lt, moq, base, estado in CATALOGO:
    serie = []
    for d in days:
        trend = 1 + 0.0015 * (d - days[0]).days
        promo = rng.random() < 0.08
        f = FIESTAS.get(fam, 1.0) if d.month == 9 and 15 <= d.day <= 19 else 1.0
        mu = base * DOW[d.weekday()] * trend * f * (1.35 if promo else 1.0)
        u = int(max(0, round(rng.normal(mu, mu * 0.15))))
        serie.append(u)
        ventas.append({
            "Fecha": d, "Producto": prod, "Familia": fam, "Unidades Vendidas": u,
            "Precio Unitario": precio, "Total Venta": u * precio, "Promocion": "Si" if promo else "No",
        })
    s90 = pd.Series(serie[-90:], index=pd.to_datetime(days[-90:]), dtype=float)
    params = {"unit_cost": float(costo or 1.0), "order_cost": 5000.0, "holding_rate": 0.25,
              "lead_time_days": lt, "moq": float(moq or 0)}
    s, S, _ = _calc_s_S(s90, params)
    qty = int(s * 0.4) if estado == "critico" else int(S * 1.2) + 1
    stock.append({"Producto": prod, "Familia": fam, "Stock Disponible": qty, "Unidad": "un"})
    esperado.append((prod, prov, estado, qty, s, S, costo, moq))

catalogo = pd.DataFrame([{
    "Producto": p, "Familia": f, "Proveedor": pr or "", "Costo Unitario": c,
    "Costo Pedido": 5000, "Lead Time (dias)": lt, "MOQ": m,
} for p, f, pr, c, _pv, lt, m, _b, _e in CATALOGO])

# Resultado esperado de "Generar ordenes": criticos con costo, agrupados por proveedor
ocs: dict = {}
for prod, prov, estado, qty, s, S, costo, moq in esperado:
    if estado == "critico" and costo is not None:
        ocs.setdefault(prov or "Sin proveedor", []).append(f"{prod} (~{max(S - qty, moq or 0):.0f} un)")
ocs_txt = [f"   - {prov}: {', '.join(items)}" for prov, items in ocs.items()]

criticos = [p for p, _, e, *_ in esperado if e == "critico"]
instrucciones = [
    ["SMARTSUPPLY - LIBRO DE PRUEBA (Distribuidora El Roble)", ""],
    ["", ""],
    ["Contenido", "Catalogo (22 productos, 6 familias, 10 proveedores), Ventas (180 dias, unidades + monto CLP + promociones) y Stock actual."],
    ["", f"Rango de ventas: {days[0]:%d/%m/%Y} a {end:%d/%m/%Y}. Incluye peak de Fiestas Patrias en BEBIDAS, SNACKS y PANADERIA."],
    ["", ""],
    ["PASO", "QUE HACER / QUE DEBERIAS VER"],
    ["0. Mi negocio", "Antes de cargar: arriba elige el negocio (o crea uno en Mi negocio). Completa RUT, ciudad y direccion: salen en el PDF de cada OC."],
    ["1. Ingesta", "Revisa que el destino (negocio y ubicacion) sea el correcto y sube este libro completo. Debe decir 'Libro con varias hojas detectado' y mostrar tres tablas: Catalogo (22), Stock (22) y Ventas (3.960 filas). Esta hoja se ignora con un aviso."],
    ["", "Las ventas traen unidades y monto: debe aparecer 'Unidades y monto CLP detectados'. Elige negocio + ubicacion y confirma."],
    ["2. Productos", "22 productos con Familia y Proveedor. GALLETAS TRITON 126G sin costo (rojo). Prueba el buscador con 'Soprole' o 'BEBIDAS'."],
    ["3. Forecast", "El selector de producto agrupa por familia (ABARROTES, BEBIDAS, ...). Corre COCA-COLA 1.5L a 14 dias: deberia notar el ciclo semanal (sabado alto, domingo bajo)."],
    ["4. Inventario", f"{len(criticos)} productos criticos (stock bajo el punto s): {', '.join(criticos)}."],
    ["", "GALLETAS TRITON 126G aparece critico pero con aviso 'falta configurar costo' (no entra en ordenes automaticas)."],
    ["5. Ordenes - generar", f"'Generar ordenes' debe crear {len(ocs)} OC, una por proveedor:"],
    *[["", t] for t in ocs_txt],
    ["6. Ordenes - buscar", "Busca por numero (ej. OC-000123 o 123), por producto (COCA-COLA, LECHE) o por proveedor (Andina). Filtra por estado y fechas."],
    ["7. Ordenes - flujo", "Abre una OC: revisa lineas y total, edita cantidades (solo pendiente), descarga el PDF, confirmala, marca en transito y luego recibida."],
    ["", "Al recibirla, el stock de cada producto sube en Inventario y deja de estar critico."],
    ["8. Ordenes - manual", "Crea una OC a mano para 'Colun': agrega MANTEQUILLA COLUN 250G x 12. El costo se completa desde el catalogo."],
    ["9. Ordenes - cancelar", "Cancela una OC pendiente: queda en estado Cancelada y el producto vuelve a poder generarse automaticamente."],
]

with pd.ExcelWriter(OUT, engine="openpyxl") as xw:
    pd.DataFrame(instrucciones).to_excel(xw, sheet_name="Instrucciones", header=False, index=False)
    catalogo.to_excel(xw, sheet_name="Catalogo", index=False)
    pd.DataFrame(ventas).to_excel(xw, sheet_name="Ventas", index=False)
    pd.DataFrame(stock).to_excel(xw, sheet_name="Stock", index=False)

wb = load_workbook(OUT)
head = PatternFill("solid", fgColor="1E293B")
for ws in wb.worksheets:
    if ws.title == "Instrucciones":
        ws.column_dimensions["A"].width = 22
        ws.column_dimensions["B"].width = 140
        ws["A1"].font = Font(bold=True, size=14)
        for row in ws.iter_rows():
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
            if row[0].value and str(row[0].value)[:1].isdigit() or row[0].value in ("PASO", "Contenido"):
                row[0].font = Font(bold=True)
        continue
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(12, min(30, max(len(str(c.value or "")) for c in col[:50]) + 2))
    ws.freeze_panes = "A2"
    hdr = {c.value: c.column_letter for c in ws[1]}
    for name, fmt in (("Fecha", "DD/MM/YYYY"), ("Precio Unitario", '"$"#,##0'), ("Total Venta", '"$"#,##0'),
                      ("Costo Unitario", '"$"#,##0'), ("Costo Pedido", '"$"#,##0')):
        if name in hdr:
            for c in ws[hdr[name]][1:]:
                c.number_format = fmt
wb.save(OUT)
print(OUT)
for t in ocs_txt:
    print(t)
